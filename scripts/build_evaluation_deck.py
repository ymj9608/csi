"""Build the editable CSI rubric deck; presentation tooling is separate from runtime.

Requires python-pptx and Pillow. Run from any directory. Layout is 1440 x 810
logical pixels; all explanatory diagrams are native PowerPoint shapes/text.
"""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import json
import re

from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs/presentations'
OUT.mkdir(parents=True, exist_ok=True)
FILE = OUT / 'CSI_Evaluation_Design.pptx'
NAVY, INK, ORANGE = '061831', '152840', 'FF6B35'
WHITE, MUTED, PALE, LINE = 'FFFFFF', '64748B', 'B9C8DD', 'D7DFEA'
LIGHT, BLUE, GREEN, PURPLE = 'F2F5FA', '365DA3', '247264', '7353A2'
FONT, EN = 'Apple SD Gothic Neo', 'Arial'
FONT_PATH = '/System/Library/Fonts/AppleSDGothicNeo.ttc'
SCALE = 2 / 3
CHECKS = []


def u(n): return Pt(n * SCALE)
def rgb(s): return RGBColor.from_string(s)


def shape(s, x, y, w, h, fill=WHITE, border=None, kind=MSO_SHAPE.RECTANGLE):
    sh = s.shapes.add_shape(kind, u(x), u(y), u(w), u(h))
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    else: sh.fill.background()
    if border:
        sh.line.color.rgb = rgb(border); sh.line.width = u(1.4)
    else: sh.line.fill.background()
    sh._element.spPr.append(OxmlElement('a:effectLst'))
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0] = .04
    return sh


def txt(s, x, y, w, h, value, size=28, color=INK, bold=False,
        align='left', font=FONT, valign='top', spacing=1.18):
    sh = s.shapes.add_textbox(u(x), u(y), u(w), u(h))
    tf = sh.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {'top': MSO_ANCHOR.TOP, 'middle': MSO_ANCHOR.MIDDLE}[valign]
    for i, line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if not i else tf.add_paragraph()
        p.text = line; p.line_spacing = spacing
        p.space_before = p.space_after = Pt(0)
        p.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}[align]
        for r in p.runs:
            r.font.name = font; r.font.size = u(size); r.font.bold = bold; r.font.color.rgb = rgb(color)
            ea = OxmlElement('a:ea'); ea.set('typeface', FONT); r._r.get_or_add_rPr().append(ea)
    # Conservative logical-pixel wrapping check; real Keynote rendering is checked separately.
    ft = ImageFont.truetype(FONT_PATH, round(size), index=0)
    n = 0
    for line in value.split('\n'):
        acc = ''; n += 1
        for ch in line:
            if ft.getlength(acc + ch) > w - 2 and acc:
                n += 1; acc = ch
            else: acc += ch
    if n * size * spacing > h + 5:
        CHECKS.append({'slide': len(P.slides), 'text': value, 'height': h,
                       'estimated': round(n * size * spacing), 'width': w})
    return sh


def arrow(s, x1, y1, x2, y2, color=MUTED, width=2, head=True, dashed=False):
    sh = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, u(x1), u(y1), u(x2), u(y2))
    sh.line.color.rgb = rgb(color); sh.line.width = u(width)
    sh._element.spPr.append(OxmlElement('a:effectLst'))
    if dashed: sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if head:
        a = OxmlElement('a:tailEnd'); a.set('type', 'triangle'); a.set('w', 'sm'); a.set('len', 'sm')
        sh._element.spPr.get_or_add_ln().append(a)
    return sh


def path(s, points, color=MUTED, dashed=False):
    for i, (a, b) in enumerate(zip(points, points[1:])):
        arrow(s, *a, *b, color, head=i == len(points) - 2, dashed=dashed)


def badge(s, label, x, y, w=220, bg=LIGHT, color=BLUE):
    shape(s, x, y, w, 36, bg)
    txt(s, x+10, y+4, w-20, 28, label, 20, color, True, valign='middle')


def number(s, n, x, y, color=ORANGE, d=36):
    shape(s, x, y, d, d, color, kind=MSO_SHAPE.OVAL)
    txt(s, x, y, d, d, str(n), 22, WHITE, True, align='center', valign='middle', font=EN)


def box(s, x, y, w, h, title, body='', accent=BLUE, size=28, body_size=24):
    shape(s, x, y, w, h, WHITE, LINE)
    shape(s, x, y, 5, h, accent)
    txt(s, x+17, y+14, w-34, 39, title, size, INK, True)
    if body: txt(s, x+17, y+60, w-34, h-66, body, body_size, MUTED)


def strip(s, text, y=699, color=INK):
    shape(s, 56, y, 1328, 49, LIGHT)
    shape(s, 56, y, 4, 49, ORANGE)
    txt(s, 76, y+9, 1288, 34, text, 24, color, True)


def evidence_table(s, x, y, widths, headers, rows, row_heights, size=20):
    """Native, editable table matching the supplied two-panel evidence reference."""
    heights = [52, *row_heights]
    sh = s.shapes.add_table(len(rows)+1, len(headers), u(x), u(y),
                            u(sum(widths)), u(sum(heights)))
    t = sh.table
    t.first_row = True; t.horz_banding = False; t.vert_banding = False
    for col, width in zip(t.columns, widths): col.width = u(width)
    for row, height in zip(t.rows, heights): row.height = u(height)
    for ri, values in enumerate([headers, *rows]):
        for ci, value in enumerate(values):
            cell = t.cell(ri, ci)
            cell.fill.solid(); cell.fill.fore_color.rgb = rgb(NAVY if ri==0 else WHITE)
            cell.margin_left = cell.margin_right = u(9)
            cell.margin_top = cell.margin_bottom = u(6)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame; tf.clear(); tf.word_wrap = True
            fs = 22 if ri==0 else size
            for pi, line in enumerate(value.split('\n')):
                p = tf.paragraphs[0] if pi==0 else tf.add_paragraph()
                p.text = line; p.space_before = p.space_after = Pt(0); p.line_spacing = u(fs*1.10)
                p.alignment = PP_ALIGN.CENTER if ri==0 or ci==len(headers)-1 and headers[-1]=='근거' else PP_ALIGN.LEFT
                for run in p.runs:
                    run.font.name = FONT; run.font.size = u(fs)
                    run.font.bold = ri==0
                    run.font.color.rgb = rgb(WHITE if ri==0 else INK)
                    ea = OxmlElement('a:ea'); ea.set('typeface', FONT); run._r.get_or_add_rPr().append(ea)
            tc = cell._tc.get_or_add_tcPr()
            for index, side in enumerate(['lnL','lnR','lnT','lnB']):
                ln = OxmlElement('a:'+side); ln.set('w',str(int(u(.75))))
                fill = OxmlElement('a:solidFill'); color = OxmlElement('a:srgbClr')
                color.set('val',LINE); fill.append(color); ln.append(fill); tc.insert(index,ln)
            font = ImageFont.truetype(FONT_PATH,round(fs),index=0)
            count = 0
            for line in value.split('\n'):
                count += 1; current = ''
                for ch in line:
                    if font.getlength(current+ch)>widths[ci]-20 and current:
                        count += 1; current = ch
                    else: current += ch
            need = count*fs*1.10
            if need>heights[ri]-12+4:
                CHECKS.append({'slide':len(P.slides),'cell':[ri,ci],'text':value,
                               'height':heights[ri]-12,'estimated':round(need)})
    return sh


def new(title, area, points, section, source, notes):
    s = P.slides.add_slide(P.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(WHITE)
    shape(s, 0, 0, 1440, 110, NAVY)
    shape(s, 56, 30, 4, 52, ORANGE)
    txt(s, 76, 17, 960, 25, section, 16, PALE, True, font=EN)
    txt(s, 76, 48, 1010, 52, title, 36, WHITE, True)
    txt(s, 1120, 27, 264, 28, area, 22, WHITE, True, align='right')
    txt(s, 1120, 60, 264, 24, f'배점 {points}점' if points else '평가 근거 / 보충 설명', 18, PALE, align='right')
    arrow(s, 56, 766, 1384, 766, LINE, 1, False)
    txt(s, 56, 780, 1245, 22, source, 14, MUTED)
    txt(s, 1320, 778, 64, 26, f'{len(P.slides):02d}', 20, MUTED, align='right', font=EN)
    s.notes_slide.notes_text_frame.text = notes + '\n\n검토 기준: 2026-10-21, 코드 e9a7048. 이번 파일은 평가 배점을 설명하며 획득 점수를 선언하지 않습니다.'
    return s


P = Presentation()
P.slide_width = Inches(13.333333); P.slide_height = Inches(7.5)
P.core_properties.title = 'CSI — 평가표 기반 Architecture · AI Agent · MCP 설계와 구현'
P.core_properties.subject = 'Capstone 평가표 73쪽의 담당 영역 60점에 대한 설계·구현·검증 대응'
P.core_properties.author = 'CSI Team'
P.core_properties.keywords = 'CSI, Architecture, ISO/IEC 25010, LangGraph, MCP, Capstone'

# 01 — Reading map for the examiner.
s = P.slides.add_slide(P.slide_layouts[6])
s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(NAVY)
shape(s, 56, 50, 4, 38, ORANGE)
txt(s, 77, 55, 1000, 32, 'CSI  /  ARCHITECTURE · AGENT · MCP', 23, PALE, True, font=EN)
shape(s, 1198, 35, 186, 71, WHITE)
s.shapes.add_picture(str(ROOT/'frontend/public/brand/csi-wordmark.png'), u(1209), u(47), width=u(164))
txt(s, 56, 129, 1300, 135, '평가 기준을\n설계와 실행 근거로 연결했습니다', 51, WHITE, True, spacing=1.1)
txt(s, 58, 296, 1270, 42, '설계 의도 → 구현 위치 → 검증 결과 순서로 확인하는 담당 영역 60점', 29, PALE)
cards = [
    ('10', '아키텍처 설계', '전체 AI 구조\nISO/IEC 25010 고려', '02–03 / 부록 11'),
    ('15', 'AI Agent 설계', 'State · Node · 분기 · 반복\n업무에 맞는 Agentic 흐름', '04 / 부록 12'),
    ('15', 'AI Agent 구현', '모델의 도구 선택과 근거\n상태 저장 · 재개 · 이벤트', '05–06 / 10'),
    ('20', 'MCP 설계·구현', '업무 조회 · 변경 · RAG\n연결 호출 · 예외 처리', '07–09 / 10'),
]
for i, (pts, title, desc, pages) in enumerate(cards):
    x = 56 + i*338
    arrow(s, x, 389, x, 690, '38506A', 1, False)
    txt(s, x+20, 399, 250, 68, pts, 54, ORANGE, True, font=EN)
    txt(s, x+102, 425, 100, 29, '점 배점', 22, PALE)
    txt(s, x+20, 486, 299, 42, title, 29, WHITE, True)
    txt(s, x+20, 551, 299, 81, desc, 24, PALE, spacing=1.3)
    txt(s, x+20, 653, 299, 29, pages, 20, WHITE, font=EN)
txt(s, 56, 752, 1328, 28, '본편 10장 + 부록 3장  ·  수치는 배점이며, 자체 평가 점수나 ISO 인증을 의미하지 않습니다.', 19, PALE)
s.notes_slide.notes_text_frame.text = (
    '발표 구성: 본편 약 8–10분을 권장합니다. 이 장은 담당 영역의 평가 항목과 뒤의 증거 위치를 연결합니다. '
    '근거: AI-Agent Capstone(+MCP).pdf 인쇄 쪽수 73, PDF 물리 페이지 74. 전체 평가표 중 담당한 네 영역만 다룹니다. '
    '아키텍처 10, AI Agent 설계 15, AI Agent 구현 15, MCP 설계·구현 20 = 배점 60. '
    'Harness·인증·대기열은 별도 배점으로 합산하지 않고 관련 설계 근거로만 언급합니다. '
    '기존 CSI.pdf의 네이비·화이트·오렌지, 선 중심 디자인을 이어갑니다. '
    '뒤 장표의 과거 실행 기록은 오늘의 재실행과 구분합니다.'
)

# 02 — Architecture in the reference's two-column evidence-table format.
s = new('아키텍처 구성과 설계 판단 근거', '아키텍처 설계', 10,
        '02. ARCHITECTURE  /  구성요소 · 책임 · 데이터 흐름',
        '평가 기준: Biz. 아이템을 위한 전체 AI 아키텍처 설계 · 코드 근거의 전체 경로는 발표자 노트 참고',
        '첨부된 공유 State/Graph 설계 표와 같은 두 열 구성입니다. 왼쪽은 구성요소·실제 기술·책임, 오른쪽은 세부 평가 관점·설계 내용·코드 근거를 연결합니다. '
        '정적 검사에서 결과 저장과 화면 갱신까지 완성된 경로를 설명합니다. 202는 접수 성공이며 최종 분석 완료를 의미하지 않습니다. '
        '결과 발행은 승인 대기보다 먼저 수행되므로 완료 보고서와 PENDING 승인은 동시에 존재할 수 있습니다. '
        '승인 요청은 Harness가 생성하고 승인 후 실제 변경은 CSI 내부 이슈 생성으로 제한됩니다. '
        'Backend 소비 재시도 소진의 DLT는 즉시 FAILED를 저장하는 전용 소비 경로가 없고, 승인 재개 실패는 기존 완료 보고서를 덮어쓰지 않습니다. '
        '근거 범례: D docs/ARCHITECTURE.md; I backend/src/main/java/com/csi/Intake.java; '
        'K backend/src/main/java/com/csi/KafkaPipeline.java; C backend/src/main/java/com/csi/ApiController.java; '
        'S backend/src/main/java/com/csi/Security.java; G agent/csi_agent/graph.py; A agent/csi_agent/app.py; '
        'M mcp_servers/business.py 및 rag.py; T tests/ 및 docs/verification.json, docs/CHANGE_REVIEWS.md. '
        '서비스 배포는 frontend/compose.yml, UI는 frontend/src/views/Console.vue, 벡터 검색은 mcp_servers/vector.py에서 확인합니다.')
txt(s,56,137,646,38,'구성요소와 책임 분리',30,INK,True)
txt(s,738,137,646,38,'아키텍처 흐름과 설계 근거',30,INK,True)
components = [
    ('검사 진입','Actions / CLI\nSemgrep · AST','정적 검사·변경 근거를 마스킹하고\nHMAC으로 Backend에 접수한다.'),
    ('웹 콘솔','Vue 3\nREST / SSE','결과 조회·검토·승인 UI를 제공한다.\nSSE 알림 후 REST로 재조회한다.'),
    ('Backend','Spring Boot\nSecurity · Outbox','인증·검사 접수·결과 검증을 맡고,\n업무 데이터와 Outbox를 저장한다.'),
    ('비동기 전달','Kafka\n검사·결과·진행·로그','검사 접수와 장시간 분석을 분리하고,\n승인 재개·결과·진행 로그를 전달한다.'),
    ('Agent','FastAPI · LangGraph\nLLM API','근거 조사·리뷰·상태 재개를 맡는다.\nRed/Blue는 한 Agent의 내부 단계다.'),
    ('MCP 서비스','business-mcp\nrag-mcp','업무는 Backend 내부 API로 접근하고,\n문서 근거는 Qdrant에서 조회한다.'),
    ('저장소','PostgreSQL\nQdrant','업무 기록·checkpoint는 PostgreSQL,\n문서 벡터는 Qdrant에 저장한다.'),
]
evidence_table(s,56,186,[121,205,320],['구분','실제 구성','담당 역할'],components,[68]*7,size=20)
architecture = [
    ('업무 목표와\n전체 AI 구조','코드 변경 → 정적 근거 → 정책·문서 대조 → 리뷰.\n실제 수집·승인 여부는 미확인 사항으로 구분한다.','D, G\nM'),
    ('비동기 처리와\n결과 정합성','검사·Outbox를 함께 저장한 뒤 202를 반환한다.\nKafka 결과를 Backend가 검증·저장하고\nSSE로 갱신을 알려 Vue가 REST로 재조회한다.','I, K'),
    ('승인과\n상태 재개','결정·Outbox 저장 후 기존 scan.request로 전달한다.\nAgent가 영속 승인 상태를 재조회하고\n같은 scanId의 checkpoint를 재개한다.','C, A\nG'),
    ('실패와\n중복 전달','멱등 처리와 유한 재시도를 적용한다.\nOutbox 최대 3회·Worker 2회·소비 초기+2회.\n위치에 따라 실패 기록 또는 DLQ/DLT로 보낸다.','I, K\nA'),
    ('인증과\n데이터 경계','세션·CSRF / HMAC / Bearer를 분리한다.\nbusiness-mcp는 업무 DB에 직접 쓰지 않는다.\n승인 후 변경은 CSI 내부 이슈 생성으로 제한한다.','S, I\nC, M'),
    ('품질과\n검증 범위','품질 9특성을 요구사항 → 설계 → 검증에 연결한다.\n샘플 통과와 성능·정확도의 정량 평가를 구분한다.','D, T'),
]
evidence_table(s,738,186,[133,441,72],['평가 관점','설계 내용과 판단 근거','근거'],architecture,[79]*6,size=20)
txt(s,56,725,1328,19,'근거  D ARCHITECTURE.md   I Intake.java   K KafkaPipeline.java   C ApiController.java   S Security.java',15,MUTED)
txt(s,56,746,1328,19,'        G graph.py   A app.py   M business.py / rag.py   T 테스트·통합 기록  ·  전체 파일 경로와 예외 사항은 발표자 노트',15,MUTED)

# 03 — Quality design and validation scope, also in editable side-by-side tables.
s = new('품질 기준을 반영한 아키텍처 설계', '아키텍처 설계', 10,
        '02. ARCHITECTURE  /  ISO/IEC 25010:2023 대응',
        '평가 기준: ISO/IEC 25010 품질평가를 고려한 아키텍처 설계 · 상세 요구사항·코드·검증은 ARCHITECTURE.md',
        '평가표는 표준의 판을 지정하지 않아 2023판의 9개 제품 품질 특성을 기준으로 연결했습니다. '
        '왼쪽에 주요 기능·실행 통제 특성 5개, 오른쪽에 운영·확장 특성 4개와 검증의 해석 범위를 배치했습니다. '
        'QA-01부터 QA-09는 docs/ARCHITECTURE.md의 요구사항 추적 ID이며, 코드 위치와 V1–V4 실행 근거에 연결됩니다. '
        'QA-01 기능 적합성, QA-02 성능 효율성, QA-03 호환성, QA-04 상호작용 능력, QA-05 신뢰성, QA-06 보안성, '
        'QA-07 유지보수성, QA-08 유연성, QA-09 안전성. '
        'V1/V2는 2026-10-21 Python 40개/Java 8개 테스트, V3는 2026-10-19 Local Demo 통합, V4는 2026-10-20 변경 리뷰·E2E 기록입니다. '
        '테스트 대역과 실제 서비스 통합은 구분합니다. 기준 개수는 정확도·커버리지 비율이 아닙니다. '
        '2011판 사용성/이식성은 2023판 상호작용 능력/유연성으로 개정되었고 안전성이 추가됐습니다. '
        '안전성은 오판과 자동 실행의 운영 피해를 줄이는 제한된 설계 관점이며 정식 위해 분석은 아닙니다. '
        '공식 출처 https://www.iso.org/standard/78176.html . ISO 인증이나 전체 품질 특성 충족 선언이 아닙니다.')
txt(s,56,137,646,38,'핵심 품질 요구와 설계',30,INK,True)
txt(s,738,137,646,38,'운영·확장 품질과 검증 범위',30,INK,True)
core_quality = [
    ('기능 적합성','실제 Finding·정적 사실·인용 ID를 검증한다.\n가짜 근거 거부와 샘플 검사를 확인했으나,\n대표 데이터셋의 오탐·미탐은 미측정이다.','QA-01'),
    ('성능 효율성','202 비동기 접수와 모델·도구 예산을 적용한다.\n반복·입력·호출 제한은 테스트했으나,\np95 지연·처리량·실제 비용은 미측정이다.','QA-02'),
    ('신뢰성','Outbox·멱등성·checkpoint·유한 재시도 적용.\n승인·중복·Kafka 장애의 실행 기록이 있다.\n강제 종료 후 복구·대규모 부하는 미검증이다.','QA-05'),
    ('보안성','인증 경계·마스킹·읽기 도구 allowlist를 둔다.\n변조·익명·CSRF 거부를 확인했으나,\n종합 공격 대응과 다중 조직 권한은 과제다.','QA-06'),
    ('안전성','HITL·Backend 승인 재검증으로 쓰기를 통제한다.\n승인 시 내부 이슈 1건·거절 시 0건을 확인했다.\n자동 패치·대상 앱 실행·운영 배포는 하지 않는다.','QA-09'),
]
other_quality = [
    ('호환성','JSON 계약·Streamable HTTP로 서비스를 잇는다.\n구버전 조회와 로컬 통합 기록은 있으나,\n다양한 클라이언트·버전 조합은 미검증이다.','QA-03'),
    ('상호작용 능력','상태·근거·불확실성을 구분해 표시한다.\n화면·SSE·모바일 E2E 기록은 있으나,\n사용자 과업·접근성 평가는 부족하다.','QA-04'),
    ('유지보수성','규칙·프롬프트·업무 API와 테스트를 분리한다.\nPython 40개·Java 8개 회귀 테스트 통과.\n변경 비용·전체 분기 커버리지는 미측정이다.','QA-07'),
    ('유연성','Compose·환경 변수로 배포 설정을 분리한다.\n로컬 빌드·기동 기록은 있으나,\n환경별 재현·다중 Worker·고가용성은 미검증이다.','QA-08'),
    ('검증의\n해석 범위','구현·샘플 검증·미측정 항목을 구분한다.\n단위 테스트 대역과 과거 라이브 통합을 나누며,\n전체 품질 달성이나 ISO 인증으로 설명하지 않는다.','V1–V4'),
]
evidence_table(s,56,186,[137,426,83],['품질 기준','설계 내용과 확인 범위','근거'],core_quality,[96]*5,size=20)
evidence_table(s,738,186,[137,426,83],['품질 기준','설계 내용과 확인 범위','근거'],other_quality,[96]*5,size=20)
txt(s,56,728,1328,19,'근거  QA-01–09: docs/ARCHITECTURE.md의 품질 요구사항·설계·코드·검증 대응표',15,MUTED)
txt(s,56,747,1328,18,'        V1·V2: 2026-10-21 단위/경계 테스트  /  V3: 2026-10-19 통합 기록  /  V4: 2026-10-20 변경 리뷰·E2E 기록',14.5,MUTED)

# 04 — State and actual conditional loop.
s = new('업무 근거의 부족 여부가 다음 행동을 결정합니다', 'AI Agent 설계', 15,
        '03. AGENT & MCP  /  STATE · GRAPH · BRANCH · LOOP',
        '근거: agent/csi_agent/graph.py State / build_graph · planner.py missing_requirements',
        '한 개 LangGraph Agent이며 Red/Blue는 내부 분석 단계입니다. State는 event, findings, changes, policy, references, comparison, '
        'action, observation, trace, toolCounts, budget, approval, decision 등을 보존합니다. '
        '초기 검사 검증 및 Finding/변경 조회는 고정된 업무 단계입니다. 현재/기준 Finding이나 변경 파일이 있으면 선택 반복으로 들어갑니다. '
        'finish_analysis 요청이 있어도 정책·POL-01 비교·규칙별 내부/외부 근거가 부족하면 종료를 거부합니다. '
        'Red에서 정적 사실을 구성하고 Blue 단계에서 LLM 리뷰와 출력 검증을 수행합니다. 리뷰 결과 발행 후 현재 Finding이 있으면 승인 중단, 없으면 종료합니다. '
        '본 그림은 실제 노드를 묶은 요약이며 전체 edge 목록은 graph.py에 있습니다. 모든 단계가 모델의 자율 선택은 아닙니다.')
badge(s,'STATE  /  기억하는 업무 정보',56,145,329)
groups=[('검사·변경','event · findings · changes'),('근거','policy · references · comparison'),('조사 과정','action · observation · trace'),('실행·승인','budget · approval · decision')]
for i,(a,b) in enumerate(groups): box(s,56,197+i*113,329,103,a,b,BLUE,25,19)
badge(s,'GRAPH  /  고정 단계와 모델 선택 분리',421,145,963)
for x,t,b in [(421,'검사 검증','validate_scan'),(756,'근거 로드','load_findings'),(1091,'조사 필요 판단','select_analysis_tools')]:
    box(s,x,201,293,100,t,b,BLUE,26,20)
arrow(s,714,251,756,251);arrow(s,1049,251,1091,251)
shape(s,421,335,963,209,LIGHT)
txt(s,438,344,890,30,'규칙 또는 변경 존재 → 모델 선택 루프',23,BLUE,True)
for x,t,b in [(438,'다음 도구 선택','plan_analysis'),(762,'인자 검증 · 실행','execute_analysis_tool'),(1086,'관측 · 근거 확인','assess_evidence')]:
    box(s,x,390,280,100,t,b,ORANGE,25,18)
arrow(s,718,440,762,440,ORANGE);arrow(s,1042,440,1086,440,ORANGE)
path(s,[(1226,490),(1226,534),(578,534),(578,490)],ORANGE)
txt(s,680,496,494,30,'근거 부족 / 도구 오류 → 재선택',22,ORANGE,True,align='center')
path(s,[(1237,301),(1237,322),(423,322),(423,380),(578,380),(578,390)],BLUE)
for x,t,b in [(421,'Red / Blue 검토','정적 사실 + LLM 리뷰 검증'),(756,'결과 발행','scan.result → Backend 저장'),(1091,'현재 Finding 분기','있음: HITL / 없음: 종료')]:
    box(s,x,584,293,100,t,b,BLUE,25,20)
arrow(s,714,634,756,634);arrow(s,1049,634,1091,634)
path(s,[(1366,440),(1400,440),(1400,576),(568,576),(568,584)],BLUE)
txt(s,855,547,440,28,'종료 요청 + 근거 충족',20,BLUE,True)
strip(s,'조사할 규칙·변경이 없으면 루프를 건너뜁니다. Red / Blue는 한 Agent 안의 분석 단계입니다.')

# 05 — Real evidence of model choice, separate from our schematic.
s = new('모델의 선택과 근거 인용을 실행 기록으로 확인합니다', 'AI Agent 구현', 15,
        '03. AGENT & MCP  /  모델 선택 → 관측 → 근거 기반 답변',
        '기록: docs/AGENT_HARNESS.md · 2026-10-19 LOCAL_DEMO · scan 3e963769…',
        '실제 실행 기록을 도식화한 장표이며 새 실행 결과나 원본 화면 캡처는 아닙니다. '
        'scanId 3e963769-b666-4d74-be45-5dee03b54769, 2026-10-19 LOCAL_DEMO, gpt-4.1-mini, COMPLETED/LLM. '
        '선택 5턴(종료 포함) + 리뷰 배치 2회로 모델 API 호출 7회였고 API usage 11,338 tokens입니다. '
        '초기 고정 MCP 조회는 이 선택 5턴과 별개입니다. finish_analysis는 MCP 도구가 아닌 로컬 제어 동작입니다. '
        '실제 trace는 get_privacy_policy, compare_policy_with_findings, search_internal_privacy_policy, '
        'search_security_guidelines, finish_analysis 순서였습니다. 순서가 이 사례에서 이렇게 나왔으며 모든 검사에서 고정된 순서라고 설명하지 않습니다. '
        '모델 출력의 존재하지 않는 사실/문서 ID와 C1 인용 누락을 검증합니다. 참조 ID가 유효해도 모든 문장의 의미적 정확성을 보장하지 않습니다.')
badge(s,'실제 LOCAL_DEMO 실행 기록',56,141,364)
txt(s,455,145,900,34,'모델 선택 5턴  ·  리뷰 배치 2회  ·  gpt-4.1-mini',28,INK,True)
trace=[('get_privacy_policy','검사 접수 시점의 정책 조회'),
       ('compare_policy_with_findings','필드를 승인·제한 목록과 대조'),
       ('search_internal_privacy_policy','사내 정책의 조항 근거 확보'),
       ('search_security_guidelines','외부 참조집의 관련 근거 확보'),
       ('finish_analysis','필수 근거 확인 후 종료 · 로컬 제어')]
for i,(name,desc) in enumerate(trace):
    y=208+i*89
    if i<4:arrow(s,76,y+35,76,y+89,LINE)
    number(s,i+1,58,y+4)
    txt(s,112,y,746,36,name,27,INK,True,font=EN)
    txt(s,112,y+40,746,34,desc,24,MUTED)
shape(s,906,207,478,445,LIGHT)
txt(s,931,229,428,40,'근거가 답변을 제한',30,INK,True)
for i,(a,b) in enumerate([('정적 사실','규칙 · 위치 · 마스킹 근거'),('정책·문서','문서 ID · 조항 · 페이지'),('검증된 리뷰','Red / Blue · 미확인 사항')]):
    y=293+i*100
    txt(s,934,y,420,34,a,26,BLUE,True);txt(s,934,y+39,420,34,b,24,MUTED)
txt(s,933,600,422,34,'허용된 ID·인용만 수락',25,ORANGE,True)
strip(s,'trace의 LLM / RULE_BASED / LLM_FALLBACK을 구분합니다. 고정 초기 조회는 모델 선택으로 세지 않습니다.')

# 06 — Stored decisions, same-topic resumption, and bounded failure paths.
s = new('보고서는 보존하고, 승인된 후속 작업만 재개합니다', 'AI Agent 구현', 15,
        '03. AGENT & MCP  /  CHECKPOINT · HITL · RECOVERY',
        '근거: app.py · graph.py · ApiController.java · KafkaPipeline.java · verification.json',
        '보고서 발행은 승인 대기 이전입니다. 완료 보고서와 승인 PENDING은 동시에 존재할 수 있습니다. '
        'interrupt 이후 PostgresSaver에 저장되며 thread_id는 scanId입니다. 사용자의 결정은 Backend가 DB/Outbox에 저장합니다. '
        'APPROVAL_DECIDED가 기존 scan.request를 타고 전달되면 Agent가 영속 승인 상태를 Backend에서 재조회하여 Command(resume)합니다. '
        '승인 시 create_security_issue를 호출하고 Backend가 다시 승인 상태를 확인해 승인별 멱등 내부 이슈를 기록합니다. 거절 시 이슈 없이 종료합니다. '
        'Outbox 최대 3회, Worker 최대 2회, Backend 소비 초기 + 2회 재시도 후 DLT입니다. 승인 재개 실패는 기존 완료 보고서를 FAILED로 덮지 않습니다. '
        '오래된 QUEUED/RUNNING 10분 초과 처리는 있지만 COMPLETED 검사의 승인 정체는 대상이 아닙니다. DLQ 자동 재처리는 없습니다. '
        '2026-10-19 실행 기록에서 DB 체크포인트와 승인/거절을 확인했고 프로세스 강제 종료 후 복구는 별도 미검증입니다.')
txt(s,56,141,1290,41,'thread_id = scanId  ·  승인 결정은 메시지 내용만 믿지 않고 DB에서 다시 읽습니다.',29,INK,True)
stages=[('보고서 발행 후 대기','검토 요청 → interrupt()\nPostgreSQL checkpoint'),
        ('검토자의 승인 / 거절','Backend 결정 + Outbox\n동일 트랜잭션 저장'),
        ('같은 토픽으로 전달','scan.request\nAPPROVAL_DECIDED'),
        ('영속 상태 확인 · 재개','승인 상태 재조회\nCommand(resume)')]
for i,(title,body) in enumerate(stages):
    x=56+i*338;number(s,i+1,x,209)
    box(s,x,259,314,153,title,body,ORANGE,25,24)
    if i<3:arrow(s,x+314,331,x+338,331,ORANGE)
box(s,56,463,644,102,'APPROVED → 내부 이슈 1건','Backend 승인 재검증 · approval별 멱등 처리',GREEN,27,24)
box(s,740,463,644,102,'REJECTED → 이슈 생성 없이 종료','코드 변경 · 외부 GitHub 이슈 · 운영 배포 없음',MUTED,27,24)
txt(s,56,600,1328,32,'실패 경로도 분리해 무한 재시도를 막습니다.',27,INK,True)
for i,(a,b) in enumerate([('Outbox','최대 3회 → QUEUED 검사 FAILED'),('Agent Worker','최대 2회 → scan.dlq'),('Backend 소비','초기 + 2회 → 원본 토픽 .DLT')]):
    x=56+i*451
    txt(s,x,644,426,29,a,23,BLUE,True);txt(s,x,681,426,31,b,23,MUTED)
txt(s,56,731,1328,25,'검사 완료 ≠ 후속 작업 승인 완료  ·  DLQ 자동 재처리와 강제 종료 복구 시험은 남은 과제',20,MUTED)

# 07 — Complete actual tool inventory, showing boundary of model access.
s = new('업무 도구와 문서 검색 도구의 책임을 나눴습니다', 'MCP 설계·구현', 20,
        '03. AGENT & MCP  /  BUSINESS-MCP · RAG-MCP',
        '근거: mcp_servers/business.py · rag.py · planner.py READ_TOOLS',
        'business-mcp 8개(읽기 6, 변경 2), rag-mcp 3개 읽기 도구와 1개 URI 템플릿 Resource를 구현했습니다. '
        '모든 도구에 설명과 타입 인자, annotation이 있습니다. 초기 context/findings/changes 조회는 고정 흐름입니다. '
        '모델의 자유 선택은 business 3개 + rag 3개 읽기 도구이며 finish_analysis는 로컬 종료 제어입니다. '
        'request_security_review는 승인 요청 레코드를 만드는 쓰기이며 승인 이전에 Harness가 수행합니다. '
        'create_security_issue만 HITL 승인과 Backend 재검증 이후 실행합니다. 두 쓰기 도구를 모두 승인 후 실행한다고 표현하지 않습니다. '
        'business-mcp는 /internal API를 경유하며 업무 DB 직접 접근은 없습니다. rag-mcp는 Qdrant 검색을 담당합니다.')
shape(s,56,148,742,502,LIGHT);shape(s,830,148,554,502,LIGHT)
txt(s,80,169,694,46,'business-mcp',34,GREEN,True,font=EN)
txt(s,80,219,694,30,'읽기 6개 + 변경 2개  /  Backend 내부 API',24,MUTED)
txt(s,80,276,162,30,'고정 초기 조회',23,INK,True)
txt(s,260,274,514,85,'get_scan_context · get_scan_findings\nget_scan_changes',23,INK,font=EN,spacing=1.45)
arrow(s,80,365,774,365,LINE,1,False)
txt(s,80,390,162,30,'모델 선택 읽기',23,INK,True)
txt(s,260,387,514,103,'get_privacy_policy · get_scan_history\ncompare_policy_with_findings',23,INK,font=EN,spacing=1.45)
arrow(s,80,488,774,488,LINE,1,False)
txt(s,80,515,162,31,'통제된 쓰기',23,ORANGE,True)
txt(s,260,513,514,115,'request_security_review  → 검토 요청\ncreate_security_issue  → 승인 후 이슈',23,INK,spacing=1.55)
txt(s,854,169,506,46,'rag-mcp',34,PURPLE,True,font=EN)
txt(s,854,219,506,30,'읽기 3개 + Resource  /  Qdrant',24,MUTED)
txt(s,854,278,506,161,'search_internal_privacy_policy\nsearch_security_guidelines\nget_guideline_details',24,INK,font=EN,spacing=1.8)
arrow(s,854,453,1360,453,LINE,1,False)
txt(s,854,483,506,34,'MCP Resource',26,PURPLE,True,font=EN)
txt(s,854,530,506,32,'security://guidelines/{rule_id}',23,INK,font=EN)
txt(s,854,581,506,42,'규칙별 실제 조항 · 원문 위치 반환',24,MUTED)
strip(s,'모델의 선택 목록 = 읽기 도구 6개 + 로컬 finish_analysis. 쓰기 도구는 고정된 승인 워크플로가 통제합니다.')

# 08 — Protocol contract, explicit error observations, and actual enforcement.
s = new('MCP 호출은 인자·권한·예산 검증을 통과해야 합니다', 'MCP 설계·구현', 20,
        '03. AGENT & MCP  /  CONTRACT · ROUTING · ERROR HANDLING',
        '근거: planner.py validate_action · graph.py tool/execute · mcp_servers/common.py · tests/test_agent_loop.py',
        'LLM은 명시된 함수 스키마를 보고 호출을 제안합니다. Harness는 도구 allowlist와 정확한 인자 키, rule_id 범위, query 길이를 검사합니다. '
        'business 도구의 scanId는 모델이 바꾸지 못하도록 Harness가 주입합니다. FastMCP Client는 Bearer로 Streamable HTTP 서버에 연결합니다. '
        '15초 시간 제한 내 호출하며 서버 오류/예산 소진은 TOOL_ERROR 관측으로 다음 선택에 반영합니다. '
        '예산 소진 도구는 선택 목록에서 제외하고 도구 이름별 반복은 3회입니다. 최대 12턴 후에는 실패합니다. '
        '모델 선택 실패는 명시적 LLM_FALLBACK이며 무제한 반복이 아닙니다. 모든 MCP 오류가 동일 상세 원인으로 관측되는 한계가 있어 원인별 분류는 향후 과제입니다. '
        'readOnlyHint/destructiveHint는 설명 메타데이터이고 접근 통제는 토큰·인자·Backend 승인 검증이 담당합니다. 내부 HTTP는 TLS가 아닙니다.')
txt(s,56,141,1328,39,'모델은 도구를 제안하고, Harness와 Backend가 실행 범위를 제한합니다.',29,INK,True)
for i,(title,body) in enumerate([('LLM 함수 선택','허용된 읽기 함수\n도구명 + 인자'),('Harness 검증','allowlist · 인자 schema\nscanId 주입 · 호출 예산'),('FastMCP 호출','Streamable HTTP\nBearer 인증'),('서비스 경계','business → 내부 API\nrag → Qdrant')]):
    x=56+i*338;box(s,x,239,314,142,title,body,BLUE,27,24)
    if i<3:arrow(s,x+314,310,x+338,310,BLUE)
txt(s,56,411,1328,38,'오류를 다음 관측으로 전달하되, 조사와 변경을 무제한 허용하지 않습니다.',27,INK,True)
for i,(title,body) in enumerate([
    ('허용되지 않은 인자','쓰기 함수·다른 scanId·잘못된 규칙\n→ 호출 제안 거부 / fallback'),
    ('MCP 오류 · 시간 초과','TOOL_ERROR를 다음 선택에 전달\n→ 대안 조사 또는 유한 실패'),
    ('반복 · 모델 장애','도구별 3회 / 선택 최대 12턴\n→ 소진 도구 제외 · 모드 표시')]):
    box(s,56+i*451,475,426,148,title,body,ORANGE,26,23)
strip(s,'annotation은 권한 검사가 아닙니다. readOnlyHint / destructiveHint + 실제 인증·인자·승인 검증을 함께 적용합니다.')

# 09 — RAG provenance and exact business-context limitations.
s = new('업무 맥락은 확인 가능한 정책·코드 사실까지 반영합니다', 'MCP 설계·구현', 20,
        '03. AGENT & MCP  /  POLICY CONTEXT · RAG PROVENANCE',
        '근거: rag/catalog.json · mcp_servers/index.py / vector.py / policy_review.py · scripts/evidence_context.py',
        '제공된 PrivacyGuard 사내 정책/외부 참조집 2개 PDF에서 조항과 실제 페이지를 추출해 47개 청크를 만들었습니다. '
        'PG-POL 28개, PG-REF 19개입니다. 1024차원 문자 2–4그램 HashingVectorizer로 벡터화하고 Qdrant cosine 검색 및 문서 유형/규칙 필터를 적용합니다. '
        '신경망 임베딩 모델을 쓰는 구조가 아닙니다. OWASP/NIST 관련 설명이 참조집에 포함될 수 있지만 공식 원문 corpus를 직접 색인한 것은 아닙니다. '
        'policy_review는 필드를 승인/제한 목록과 대조합니다. AST는 지원 클래스와 명시적 생성자/db.add 및 실습 조건을 관찰합니다. '
        '변경 리뷰는 C1 변경 사실을 추가합니다. 수집 목적·보유 기간·별도 승인 기록·런타임 로그 자동 연계는 미구현입니다. '
        'match_business_context 도구는 없으며 이를 구현된 기능처럼 설명하지 않습니다. 근거 ID와 실제 페이지를 반환하더라도 법적 위반이나 실제 수집을 확정하지 않습니다.')
flows=[('정책 PDF 2종','사내 정책 / 외부 참조집'),('조항 · 페이지 추출','catalog.json · 47청크'),('로컬 벡터화','1024차원 · 문자 n-gram'),('Qdrant 검색','유형·규칙 필터 / cosine'),('출처 반환','문서 ID · 조항 · 페이지')]
for i,(title,body) in enumerate(flows):
    x=56+i*270;box(s,x,196,248,123,title,body,PURPLE,25,21)
    if i<4:arrow(s,x+248,257,x+270,257,PURPLE)
txt(s,56,349,1328,35,'필드가 존재한다는 사실과, 업무상 수집이 승인됐다는 판단을 구분합니다.',29,INK,True)
columns=[('반영된 업무 근거','정책의 승인·제한 목록\n필드 역할 · 정적 저장 호출\n실습 모드 조건 · 코드 변경 사실',GREEN),
         ('아직 연결되지 않은 맥락','실제 수집 목적 · 보유 기간\n별도 승인 기록 · 운영 설정\n런타임 업무 로그',MUTED),
         ('결과에 드러내는 한계','INSUFFICIENT_EVIDENCE\n별도 승인 기록 미확인\n런타임 검증 미수행',ORANGE)]
for i,(title,body,col) in enumerate(columns):box(s,56+i*451,419,426,220,title,body,col,27,24)
strip(s,'RAG 출처는 제공된 사내 정책·외부 참조집입니다. OWASP 공식 원문 전체나 실제 업무 승인 시스템과 구분합니다.')

# 10 — Evidence summary, precise dates and no inflated score claims.
s = new('설계 문서·코드·실행 결과를 함께 제시합니다', '담당 영역 종합', 60,
        '04. EVIDENCE  /  무엇을 어디까지 확인했는가',
        '근거: ARCHITECTURE.md V1–V4 · AGENT_HARNESS.md · CHANGE_REVIEWS.md · verification.json',
        '이 장은 배점 60점 영역의 검증 안내이지 획득 점수표가 아닙니다. 2026-10-21 동일 코드 e9a7048에서 Python 40개, Java 8개가 통과했습니다. '
        'Python 테스트는 Semgrep fixture 실제 실행과 모델/MCP 대역 테스트가 혼합되어 있습니다. Java 8개는 입력/출력 경계 단위 테스트입니다. '
        '2026-10-19의 HITL/장애/인증 기록과 2026-10-20의 실제 LLM 변경 리뷰/Chrome E2E 2개/Vue 빌드 기록은 과거 Local Demo입니다. '
        '오늘 새 LLM·GitHub Push·장애 주입을 실행한 것으로 표현하지 않습니다. 테스트 개수를 탐지 정확도나 전체 서비스 커버리지 비율로 해석하지 않습니다. '
        '단위/통합 검증은 있으며 대표 정확도 데이터셋, p95·비용, 부하·장애 복구 평가는 추가 과제입니다.')
for i,(n,desc) in enumerate([('40','Python 테스트 통과'),('8','Java 테스트 통과')]):
    x=56+i*380;txt(s,x,144,130,76,n,59,BLUE,True,font=EN)
    txt(s,x+138,162,235,60,desc,25,INK,True)
txt(s,874,156,510,65,'2026-10-21 · 로컬 코드 검증\n라이브 전체 시스템 재실행과 구분',24,MUTED)
data=[('아키텍처 설계','책임·승인·실패 경로 + 품질 추적표','QA-01–09 / Outbox·멱등 처리'),
      ('AI Agent 설계','State·조건 분기·선택 반복의 코드 대응','graph.py / missing_requirements'),
      ('AI Agent 구현','실제 모델 선택·근거·승인 재개 기록','2026-10-19 LLM trace / HITL'),
      ('MCP 설계·구현','실제 연결·도구 설명·권한·예외 처리','MCP Resource / 거부·재선택 시험')]
xs=[56,342,942];ws=[286,600,442]
for x,w,h in zip(xs,ws,['평가 영역','확인할 증거','대표 위치 / 기록']):
    shape(s,x,267,w,43,NAVY);txt(s,x+14,275,w-28,31,h,23,WHITE,True)
for i,row in enumerate(data):
    y=310+i*80;shape(s,56,y,1328,80,LIGHT if i%2==0 else WHITE)
    for j,(x,w,t) in enumerate(zip(xs,ws,row)):
        txt(s,x+14,y+21,w-28,48,t,25 if j==0 else 23,INK if j<2 else MUTED,j==0)
strip(s,'남은 검증: 오탐·미탐 / p95 지연·비용 / 사용자 평가 / 프로세스 강제 종료·부하·복구')

# 11 — Complete model, terminology and primary reference in presenter notes.
s = new('부록 A. ISO/IEC 25010의 9개 품질 특성 대응', '아키텍처 설계', 10,
        'APPENDIX  /  품질 요구사항 ↔ 설계 ↔ 검증 범위',
        '공식 기준: ISO/IEC 25010:2023 · https://www.iso.org/standard/78176.html · 상세: ARCHITECTURE.md',
        'QA-01부터 QA-09까지는 CSI가 부여한 추적 ID입니다. 2011판의 Usability는 2023판 Interaction capability, Portability는 Flexibility로 개정되었고 Safety가 추가됐습니다. '
        '세부 특성이 일대일 동일하다고 해석하지 않습니다. ISO 공식 페이지: https://www.iso.org/standard/78176.html ; '
        '이전판 https://www.iso.org/standard/35733.html ; 개정 내용: https://gso-sims-preview-doc-aws.s3-eu-west-1.amazonaws.com/gso-iso-iec-25010-2025-en.html . '
        '요구사항·코드 경로·검증 V1–V4 및 미실행 합격 기준 제안은 docs/ARCHITECTURE.md에 정리했습니다. '
        '접수 p95 ≤ 1초 등 제안 목표는 아직 실측하지 않았으며 달성 실적으로 발표하지 않습니다.')
txt(s,56,137,1328,39,'모든 항목을 문서화하되, 구현된 통제와 아직 측정하지 않은 품질을 구분합니다.',27,INK,True)
xs=[56,160,419,966];ws=[104,259,547,418]
for x,w,t in zip(xs,ws,['ID','품질 특성','CSI의 설계 결정','검증 상태 / 남은 범위']):
    shape(s,x,194,w,40,NAVY);txt(s,x+12,201,w-24,29,t,22,WHITE,True)
data=[('QA-01','기능 적합성','정적 근거 · 인용 ID · 변경 연결 검증','샘플 통과 / 대표 정확도 미측정'),
      ('QA-02','성능 효율성','비동기 접수 · 모델/도구 예산 제한','제한 테스트 / p95·비용 미측정'),
      ('QA-03','호환성','JSON 계약 · Streamable HTTP · 구버전 조회','로컬 통합 / 호환 행렬 미검증'),
      ('QA-04','상호작용 능력','상태·근거·불확실성 · SSE 후 재조회','화면 E2E / 접근성 평가 부족'),
      ('QA-05','신뢰성','Outbox · 멱등성 · checkpoint · 유한 재시도','승인·장애 기록 / 복구 시험 부족'),
      ('QA-06','보안성','세션 / HMAC / Bearer · 마스킹 · allowlist','거부 시험 / 종합 공격 평가 부족'),
      ('QA-07','유지보수성','규칙·프롬프트·API 분리 · 회귀 테스트','48개 테스트 / 변경 비용 미측정'),
      ('QA-08','유연성','Compose · 환경 변수 · 서비스별 배포','로컬 기동 / 환경·확장 시험 부족'),
      ('QA-09','안전성','HITL · 승인 재검증 · 임의 실행 제한','실행 통제 / 정식 위해 분석 없음')]
for i,row in enumerate(data):
    y=234+i*48;shape(s,56,y,1328,48,LIGHT if i%2==0 else WHITE)
    for j,(x,w,t) in enumerate(zip(xs,ws,row)):
        txt(s,x+12,y+10,w-24,36,t,20 if j==0 else 22,MUTED if j==3 else INK,j==1)
strip(s,'2011판 사용성 → 2023판 상호작용 능력 / 이식성 → 유연성 / 안전성 추가. 적합성 인증을 뜻하지 않습니다.')

# 12 — Exact evidence pointers suitable for a Q&A deep dive.
s = new('부록 B. 평가 문구에서 코드와 테스트로 바로 이동', '검증 근거', 0,
        'APPENDIX  /  SOURCE & TEST TRACEABILITY',
        '검증 시점: V1·V2 2026-10-21 / V3 2026-10-19 / V4 2026-10-20 · 상세는 ARCHITECTURE.md',
        '교수님이 실제 구현 여부를 물으시면 이 표의 파일과 테스트를 보여주세요. '
        '주요 테스트: test_dynamic_order_observation_replan_and_approval_resume, '
        'test_real_planner_contract_rejects_write_and_cross_scan, test_mcp_annotations_and_mutations_are_not_model_tools, '
        'test_context_bound_and_shared_documents_with_300_findings, test_budget_caps_context_calls_and_deadline, '
        'test_llm_cannot_create_findings_or_citations, test_change_context_flows_through_mcp_graph_and_result. '
        'Spring IntakeTest/EvidenceContextTest/ChangeSetTest는 입력·근거 경계를 확인합니다. '
        '메모리 체크포인터를 사용하는 단위 테스트와 PostgreSQL을 쓰는 과거 실제 통합 증거는 분리해 설명합니다. '
        '실행 명령: 저장소 루트에서 SSL_CERT_FILE=/etc/ssl/cert.pem .venv/bin/python -m pytest tests -q ; mvn -o -f backend/pom.xml test -q. '
        '브라우저 E2E는 실제 완료 데이터와 로컬 서버가 필요하며 변경 비교용 scanId 환경 변수 없이는 건너뛸 수 있습니다.')
xs=[56,334,885];ws=[278,551,499]
for x,w,t in zip(xs,ws,['평가 문구','구현의 시작점','검증의 시작점']):
    shape(s,x,155,w,45,NAVY);txt(s,x+16,166,w-32,30,t,24,WHITE,True)
rows=[('전체 AI 구조 · 품질 고려','docs/ARCHITECTURE.md\nIntake.java · KafkaPipeline.java','V1·V2 단위 검증\nV3 Outbox·중복·장애 기록'),
      ('State · 분기 · 반복 설계','agent/csi_agent/graph.py\nplanner.py · harness.py','tests/test_agent_loop.py\n종료 거절 · 오류 후 재선택'),
      ('모델 선택 · 근거 · 재개','planner.py · analysis.py · app.py\nPostgresSaver / Command(resume)','test_review_quality.py\nAGENT_HARNESS.md / V3 HITL'),
      ('MCP 연결 · 변경 · 예외','mcp_servers/business.py · rag.py\ncommon.py · graph.py tool()','annotation · 쓰기 제안 거부\n실제 Resource / 토큰 거부 기록')]
for i,row in enumerate(rows):
    y=200+i*110;shape(s,56,y,1328,110,LIGHT if i%2==0 else WHITE)
    for j,(x,w,t) in enumerate(zip(xs,ws,row)):
        txt(s,x+16,y+24,w-32,80,t,24 if j==0 else 22,INK if j<2 else MUTED,j==0,spacing=1.4)
strip(s,'테스트 대역의 통과와 라이브 LLM·MCP·DB 통합의 증거를 분리하여 보여 주세요.')

# 13 — Demo guidance, no fabricated screenshots or claimed new runs.
s = new('부록 C. 화면에서 평가 근거를 확인하는 시연 순서', '발표 · 질의응답', 0,
        'APPENDIX  /  DEMO CHECKLIST',
        '화면 경로: 검사 이력 → 검사 상세 → 취약점 상세 → 실행 로그 → 승인 대기 · 실제 데이터 필요',
        '실제 LLM 모드로 완료되고 도구 trace 및 인용을 가진 Local Demo 검사를 준비합니다. '
        '승인 시연에는 아직 PENDING인 별도 검토 요청을 사용하고 중복 클릭으로 새 승인 요청을 만들지 않습니다. '
        '이 자료 작성 중 새로운 검사나 승인 작업을 실행하지 않았습니다. '
        '확인 순서: metadata 동일성, 선택 trace의 모드·턴·관측 상태, 정적 사실/문서 ID 및 페이지, '
        '동일 scanId의 model_decision/mcp_start/mcp_complete/approval_interrupt/finalize 로그, 승인 결정 재개 후 내부 이슈. '
        '보고서에 COMPLETED가 나왔다고 안전 판정이나 승인 완료가 아닙니다. '
        '발표 질문: 왜 Agent인가? 모델이 관측을 받아 다음 읽기 도구를 선택하고 종료 조건을 검사하는 루프가 있다. '
        '왜 MCP인가? 명시적 도구 계약으로 업무 API와 근거 검색을 분리하여 연결한다. '
        '업무 맥락의 한계는? 정책/정적 사실까지이며 실제 수집 목적과 승인 기록 자동 검증은 없다.')
txt(s,56,142,1328,41,'실행 결과의 “있음”보다, 같은 scanId에서 근거와 상태가 이어지는지 보여 줍니다.',28,INK,True)
steps=[('검사 상세','Repository · SHA\n분석 모드 · 정책 버전'),('취약점 상세','Agent 도구 선택 trace\n턴 · 관측 · LLM 여부'),('근거 · 리뷰','정적 사실 / 인용 조항\nRed·Blue / 미확인 사항'),('실행 로그','MCP 시작·완료\n선택·예산·보고서 저장'),('승인 후 재개','동일 scanId 이어가기\n승인 이슈 / 거절 종료')]
for i,(title,body) in enumerate(steps):
    x=56+i*270;number(s,i+1,x,220)
    box(s,x,276,248,159,title,body,ORANGE,26,23)
    if i<4:arrow(s,x+248,353,x+270,353,ORANGE)
txt(s,56,478,1328,39,'질문을 받았을 때 사용할 설명',29,INK,True)
qa=[('왜 Agent인가?','관측 결과로 다음 도구를 선택하고\n필수 근거가 있어야 종료합니다.'),
    ('왜 MCP인가?','업무 API와 문서 검색을\n명시적 도구 계약으로 분리합니다.'),
    ('어디까지 검증했나?','정적 사실·정책·승인 흐름까지.\n실제 업무 목적과 운영 규모는 과제입니다.')]
for i,(a,b) in enumerate(qa):box(s,56+i*451,537,426,133,a,b,BLUE,26,23)
strip(s,'시연 전 조건: 실제 완료 데이터·LLM trace·참조 문서 / HITL 시연에는 PENDING 요청 필요')


def fix_package(file):
    """Add explicit notes master linkage required by Keynote; preserve schema order."""
    ns={'p':'http://schemas.openxmlformats.org/presentationml/2006/main',
        'a':'http://schemas.openxmlformats.org/drawingml/2006/main',
        'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
    tmp=file.with_suffix('.tmp.pptx')
    with ZipFile(file) as zin, ZipFile(tmp,'w',ZIP_DEFLATED) as zout:
        rels=etree.fromstring(zin.read('ppt/_rels/presentation.xml.rels'))
        notes=[r.get('Id') for r in rels if r.get('Type','').endswith('/notesMaster')]
        for info in zin.infolist():
            data=zin.read(info.filename)
            if info.filename.endswith('.xml'):
                root=etree.fromstring(data)
                if info.filename=='ppt/presentation.xml':
                    if notes and root.find('p:notesMasterIdLst',ns) is None:
                        masters=etree.Element('{'+ns['p']+'}notesMasterIdLst')
                        for rid in notes:
                            el=etree.SubElement(masters,'{'+ns['p']+'}notesMasterId'); el.set('{'+ns['r']+'}id',rid)
                        master=root.find('p:sldMasterIdLst',ns)
                        root.insert(list(root).index(master)+1 if master is not None else 0,masters)
                    root.find('p:sldSz',ns).set('type','custom')
                for ln in root.xpath('//a:ln[a:headEnd and a:tailEnd]',namespaces=ns):
                    head=ln.find('a:headEnd',ns);tail=ln.find('a:tailEnd',ns)
                    if list(ln).index(head)>list(ln).index(tail):ln.remove(head);ln.insert(list(ln).index(tail),head)
                data=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
            zout.writestr(info,data)
    tmp.replace(file)


if __name__ == '__main__':
    P.save(FILE);fix_package(FILE)
    assert len(P.slides)==13
    for si, slide in enumerate(P.slides,1):
        for sh in slide.shapes:
            assert sh.left>=0 and sh.top>=0 and sh.left+sh.width<=P.slide_width+u(1) and sh.top+sh.height<=P.slide_height+u(1), (si,sh.name)
        assert slide.notes_slide.notes_text_frame.text.strip(), si
    print(json.dumps({'file':str(FILE),'slides':len(P.slides),'layout_warnings':CHECKS},ensure_ascii=False,indent=2))
